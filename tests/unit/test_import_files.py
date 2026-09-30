"""Excel aktarımı — dosya kontrolü, okuma güvenliği, şablon ve geçici depo (docs/11, docs/09 §3)."""

import sys
import uuid
import zipfile
from collections.abc import Sequence
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from defusedxml import DTDForbidden, EntitiesForbidden, ExternalReferenceForbidden

from site_yonetim import cli
from site_yonetim.domain.imports.unit_validator import Column, ImportFileError
from site_yonetim.services import imports as svc
from tests.conftest import SettingsFactory

NOW = datetime(2026, 9, 30, 12, 0, tzinfo=UTC)
HEADER = [c.value for c in Column]


def workbook_bytes(*sheets: tuple[str, Sequence[Sequence[object]]]) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.worksheets[0])
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def unit_row(number: str) -> list[object]:
    values = {Column.BLOCK: "A", Column.NUMBER: number, Column.OWNER_FIRST: "Ali"}
    values[Column.OWNER_LAST] = "Can"
    return [values.get(c) for c in Column]


def rewrite_entry(data: bytes, name: str, transform: object) -> bytes:
    source = zipfile.ZipFile(BytesIO(data))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == name:
                content = transform(content)  # type: ignore[operator]
            target.writestr(item.filename, content)
    return out.getvalue()


def expect_file_error(code: str, data: bytes) -> ImportFileError:
    with pytest.raises(ImportFileError) as caught:
        svc.read_and_validate(data)
    assert caught.value.code == code
    return caught.value


# --- yükleme kontrolü ------------------------------------------------------------


def test_upload_needs_xlsx_extension() -> None:
    data = workbook_bytes(("Daireler", [HEADER]))
    for name in ("daireler.xls", "daireler.csv", None, "daireler.xlsx.exe"):
        with pytest.raises(ImportFileError) as caught:
            svc.check_upload(name, data)
        assert caught.value.code == "invalid_file_type"
    svc.check_upload("DAİRELER.XLSX", data)


def test_upload_is_limited_to_5_mb() -> None:
    with pytest.raises(ImportFileError) as caught:
        svc.check_upload("a.xlsx", svc.ZIP_SIGNATURE + b"0" * svc.MAX_UPLOAD_BYTES)
    assert caught.value.code == "file_too_large"


@pytest.mark.parametrize("content", [b"%PDF-1.4 <script>", b"Blok;Daire No\nA;1", b""])
def test_upload_content_signature_must_be_zip(content: bytes) -> None:
    with pytest.raises(ImportFileError) as caught:
        svc.check_upload("sahte.xlsx", content)
    assert caught.value.code == "invalid_file_content"


# --- okuma güvenliği --------------------------------------------------------------


def test_openpyxl_parses_xml_with_defusedxml() -> None:
    from openpyxl import DEFUSEDXML

    assert DEFUSEDXML is True


def test_xml_entities_are_rejected() -> None:
    """XXE: dış varlık tanımı içeren XML okunmaz (dosya sisteminden veri sızmaz)."""
    data = workbook_bytes(("Daireler", [HEADER, unit_row("1")]))
    doctype = b'<!DOCTYPE x [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'

    def inject(content: bytes) -> bytes:
        return doctype + content.replace(b">Ali<", b">&xxe;<")

    error = expect_file_error(
        "invalid_file_content", rewrite_entry(data, "xl/worksheets/sheet1.xml", inject)
    )
    chain: list[BaseException] = []
    cause: BaseException | None = error
    while cause is not None:
        chain.append(cause)
        cause = cause.__cause__ or cause.__context__
    forbidden = (DTDForbidden, EntitiesForbidden, ExternalReferenceForbidden)
    assert any(isinstance(e, forbidden) for e in chain), chain


def test_zip_bomb_is_rejected_before_parsing() -> None:
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/worksheets/sheet1.xml", b"\0" * (svc.MAX_UNCOMPRESSED_BYTES + 1))
    data = out.getvalue()
    assert len(data) < svc.MAX_UPLOAD_BYTES  # sıkıştırılmışı küçük
    expect_file_error("file_too_complex", data)


def test_too_many_zip_entries_is_rejected() -> None:
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
        for index in range(svc.MAX_ZIP_ENTRIES):
            archive.writestr(f"x/{index}.xml", "")
    expect_file_error("file_too_complex", out.getvalue())


def test_zip_that_is_not_a_workbook_is_rejected() -> None:
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("word/document.xml", "<w/>")
    expect_file_error("invalid_file_content", out.getvalue())


def test_corrupt_zip_is_rejected() -> None:
    expect_file_error("invalid_file_content", svc.ZIP_SIGNATURE + b"bozuk" * 100)


def test_corrupt_sheet_xml_is_rejected() -> None:
    data = workbook_bytes(("Daireler", [HEADER, unit_row("1")]))
    broken = rewrite_entry(data, "xl/worksheets/sheet1.xml", lambda c: c[: len(c) // 2])
    expect_file_error("invalid_file_content", broken)


def test_formulas_are_not_evaluated() -> None:
    """Yalnız saklanmış değer okunur; openpyxl ile yazılmış formülün değeri yoktur → boş."""
    data = workbook_bytes(("Daireler", [HEADER, unit_row("=1+1")]))
    result = svc.read_and_validate(data)
    assert result.importable_count == 0
    assert result.issues[0].message == "Daire No boş olamaz."


def test_daireler_sheet_is_preferred() -> None:
    data = workbook_bytes(
        ("Açıklama", [["okuyun"]]), ("Daireler", [HEADER, unit_row("1"), unit_row("2")])
    )
    assert svc.read_and_validate(data).importable_count == 2


def test_first_sheet_is_used_without_daireler() -> None:
    data = workbook_bytes(("Sayfa1", [HEADER, unit_row("7")]))
    assert [r.number for r in svc.read_and_validate(data).rows] == ["7"]


def test_missing_columns_is_a_file_error() -> None:
    expect_file_error("missing_columns", workbook_bytes(("Sayfa1", [["Blok"], ["A"]])))


# --- şablon -------------------------------------------------------------------------


def test_template_reads_back_cleanly() -> None:
    data = svc.build_template()
    svc.check_upload("daire-aktarim-sablonu.xlsx", data)
    result = svc.read_and_validate(data)
    assert result.issues == []
    assert [r.display_name for r in result.rows] == ["A-1", "A-2", "B-Z1"]
    second = result.rows[1]
    assert second.owner.phone == "+905339876543"
    assert second.tenant is not None
    assert second.tenant.last_name == "DEMİR"
    assert result.rows[2].usage.value == "commercial"


def test_template_layout() -> None:
    workbook = openpyxl.load_workbook(BytesIO(svc.build_template()))
    assert workbook.sheetnames == ["Daireler", "Açıklama"]
    sheet = workbook["Daireler"]
    assert [c.value for c in sheet[1]] == HEADER
    assert all(c.font.bold for c in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.max_row == 4  # başlık + 3 örnek


# --- geçici depo -------------------------------------------------------------------


@pytest.fixture
def store(tmp_path: Path) -> svc.ImportStore:
    return svc.ImportStore(tmp_path / "imports")


SITE, USER = uuid.uuid7(), uuid.uuid7()


def test_saved_import_is_claimed_once(store: svc.ImportStore) -> None:
    pending = store.save(SITE, USER, b"veri", NOW)
    assert pending.expires_at == NOW + timedelta(hours=6)
    [path] = list(store.root.glob("*/*/*"))
    assert path.name == f"{pending.import_id.hex}.xlsx"  # kullanıcının dosya adı diske girmez
    if sys.platform != "win32":
        assert path.stat().st_mode & 0o777 == 0o600

    with store.claim(SITE, USER, pending.import_id, NOW) as data:
        assert data == b"veri"
    with store.claim(SITE, USER, pending.import_id, NOW) as again:
        assert again is None
    assert list(store.root.glob("*/*/*")) == []


def test_import_is_bound_to_uploader_and_site(store: svc.ImportStore) -> None:
    pending = store.save(SITE, USER, b"veri", NOW)
    with store.claim(SITE, uuid.uuid7(), pending.import_id, NOW) as other_user:
        assert other_user is None
    with store.claim(uuid.uuid7(), USER, pending.import_id, NOW) as other_site:
        assert other_site is None
    with store.claim(SITE, USER, uuid.uuid4(), NOW) as unknown:
        assert unknown is None


def test_failed_confirm_puts_the_file_back(store: svc.ImportStore) -> None:
    pending = store.save(SITE, USER, b"veri", NOW)
    with pytest.raises(RuntimeError), store.claim(SITE, USER, pending.import_id, NOW):
        raise RuntimeError
    with store.claim(SITE, USER, pending.import_id, NOW) as data:
        assert data == b"veri"


def test_expired_import_cannot_be_claimed(store: svc.ImportStore) -> None:
    pending = store.save(SITE, USER, b"veri", NOW)
    later = NOW + timedelta(hours=6)
    with store.claim(SITE, USER, pending.import_id, later) as data:
        assert data is None
    assert list(store.root.glob("*/*/*")) == []


def test_purge_removes_only_expired_files(store: svc.ImportStore) -> None:
    assert store.purge_expired(NOW) == 0  # kök henüz yok
    store.save(SITE, USER, b"eski", NOW - timedelta(hours=7))
    fresh = store.save(SITE, USER, b"yeni", NOW - timedelta(hours=1))  # kaydederken eskiyi siler
    assert [p.name for p in store.root.glob("*/*/*")] == [f"{fresh.import_id.hex}.xlsx"]
    assert store.purge_expired(NOW + timedelta(hours=5)) == 1


def test_purge_imports_command(
    tmp_path: Path,
    make_settings: SettingsFactory,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    settings = make_settings(import_storage_dir=tmp_path)
    monkeypatch.setattr(cli, "get_settings", lambda: settings)
    old = datetime.now(UTC) - timedelta(hours=7)
    svc.ImportStore(tmp_path).save(SITE, USER, b"eski", old)

    assert cli.main(["purge-imports"]) == 0
    assert "1 süresi dolmuş aktarım dosyası silindi." in capsys.readouterr().out
    assert list(tmp_path.glob("*/*/*")) == []


def test_default_import_dir_is_under_temp(make_settings: SettingsFactory) -> None:
    assert make_settings().import_dir.name == "site-yonetim-imports"
    assert make_settings(import_storage_dir="/srv/imports").import_dir == Path("/srv/imports")
