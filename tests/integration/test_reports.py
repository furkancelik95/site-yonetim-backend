"""Gelir–gider raporu, aidat tahsilat özeti, Excel çıktıları — docs/04 §12, docs/06 §2.7–2.9.

Trello "BE · Dilim 6 · Gider ve rapor" — bitti ölçütü: gider eklenince özet uçları güncel
toplamı döner. Kurgu: `finance_world` (tek daire, aylık 1.000 aidat), bugün 20.06.2026.
"""

from io import BytesIO
from pathlib import Path

import httpx2
import openpyxl
import pytest
from fastapi import FastAPI

from site_yonetim.api.v1.expenses import get_file_store
from site_yonetim.services.files import FileStore
from tests.integration.finance_world import World, finalized_plan, post_run
from tests.integration.helpers import add_site_membership, create_user, login_headers


@pytest.fixture
async def report(world: World, api_app: FastAPI, tmp_path: Path) -> World:
    api_app.dependency_overrides[get_file_store] = lambda: FileStore(tmp_path)
    await finalized_plan(world)
    await post_run(world)  # Haziran: 1.000
    accounts = {a["name"]: a["id"] for a in (await world.get("/cash-accounts")).json()["items"]}
    world.lookups.update({"bank": accounts["Banka Hesabı"], "cash": accounts["Kasa"]})
    return world


async def expense(
    w: World, amount: str, day: str, *, paid: bool = False, description: str = "Elektrik"
) -> str:
    form = {
        "expense_category_id": w.lookups["category"],
        "description": description,
        "amount": amount,
        "date": day,
        "paid": "true" if paid else "false",
    }
    if paid:
        form |= {"paid_on": day, "cash_account_id": w.lookups["bank"]}
    response = await w.api.post(w.url("/expenses"), data=form, headers=w.headers)
    assert response.status_code == 201, response.text
    return str(response.json()["data"]["id"])


async def manual(w: World, amount: str, direction: str = "in") -> str:
    response = await w.post(
        "/cash-movements",
        {"cash_account_id": w.lookups["cash"], "date": "2026-06-15", "direction": direction,
         "amount": amount, "description": "Mevduat faizi"},
    )  # fmt: skip
    assert response.status_code == 201, response.text
    return str(response.json()["data"]["id"])


async def test_gider_eklenince_ozet_guncel_toplami_doner(report: World) -> None:
    before = (await report.get("/reports/income-expense", params={"year": 2026})).json()
    assert before["total_expense"] == "0.00"
    await expense(report, "400.00", "2026-06-10", paid=True)
    after = (await report.get("/reports/income-expense", params={"year": 2026})).json()
    assert after["total_expense"] == "400.00"
    assert after["months"][5]["expense"] == "400.00"
    summary = (await report.get("/expenses", params={"year": 2026})).json()["summary"]
    assert summary["total"] == "400.00"


async def test_income_expense_report(report: World) -> None:
    paid = await report.post(
        "/payments",
        {"ledger_account_id": report.account_id, "amount": "600.00", "date": "2026-06-18",
         "method": "cash", "cash_account_id": report.lookups["cash"]},
    )  # fmt: skip
    assert paid.status_code == 201, paid.text
    await expense(report, "400.00", "2026-06-10", paid=True)
    await expense(report, "250.00", "2026-05-12")
    reversed_id = await expense(report, "300.00", "2026-06-11", paid=True)
    await report.post(f"/expenses/{reversed_id}/reverse", {"reason": "Mükerrer"})
    await manual(report, "100.00")
    wrong = await manual(report, "50.00")
    await report.post(f"/cash-movements/{wrong}/reverse", {"reason": "Yanlış giriş"})

    body = (await report.get("/reports/income-expense", params={"year": 2026})).json()
    assert len(body["months"]) == 12
    june, may = body["months"][5], body["months"][4]
    assert (june["period"], june["income"], june["expense"], june["difference"]) == (
        "06/2026",
        "700.00",
        "400.00",
        "300.00",
    )  # fmt: skip — gelir: tahsilat 600 + faiz 100 (ters kaydı alınan 50 düşer)
    assert (may["income"], may["expense"]) == ("0.00", "250.00")
    assert (body["total_income"], body["total_expense"], body["difference"]) == (
        "700.00", "650.00", "50.00"
    )  # fmt: skip
    [category] = body["categories"]
    assert (category["amount"], category["count"], category["share"]) == ("650.00", 2, "100.00")
    assert body["budget_plan"]["status"] == "finalized"
    [line] = body["budget"]
    assert (line["budgeted"], line["actual"], line["usage"], line["is_over"]) == (
        "12000.00", "650.00", "5.42", False
    )  # fmt: skip
    # Banka: −400 (gider) −300 +300 (geri alma) · Kasa: +600 +100 +50 −50
    assert body["cash_balance"] == "300.00"
    assert body["years"] == [2026]
    empty = (await report.get("/reports/income-expense", params={"year": 2025})).json()
    assert (empty["total_income"], empty["budget_plan"], empty["budget"]) == ("0.00", None, [])


async def test_collection_summary(report: World) -> None:
    await post_run(report, "2026-05-01")
    await report.post(
        "/payments",
        {
            "ledger_account_id": report.account_id,
            "amount": "1600.00",
            "date": "2026-06-18",
            "method": "cash",
        },
    )  # FIFO: Mayıs 1.000 kapanır, Haziran'a 600
    body = (await report.get("/reports/collections", params={"year": 2026})).json()
    assert [(p["period"], p["charged"], p["collected"], p["rate"]) for p in body["periods"]] == [
        ("05/2026", "1000.00", "1000.00", "100.00"),
        ("06/2026", "1000.00", "600.00", "60.00"),
    ]
    assert body["periods"][1]["outstanding"] == "400.00"
    assert (body["total_charged"], body["total_collected"], body["rate"]) == (
        "2000.00",
        "1600.00",
        "80.00",
    )


async def test_reversed_run_is_not_in_collections(report: World) -> None:
    run_id = (await report.get("/charge-runs")).json()["items"][0]["id"]
    await report.post(f"/charge-runs/{run_id}/reverse", {"reason": "Hata"})
    body = (await report.get("/reports/collections", params={"year": 2026})).json()
    assert body["periods"] == []
    assert body["rate"] == "0.00"


# --- Excel -------------------------------------------------------------------------------


def workbook(response: httpx2.Response) -> openpyxl.Workbook:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"].startswith("attachment;")
    return openpyxl.load_workbook(BytesIO(response.content))


async def test_report_export_has_three_sheets(report: World) -> None:
    await expense(report, "400.00", "2026-06-10", paid=True)
    book = workbook(await report.get("/reports/income-expense/export.xlsx", params={"year": 2026}))
    assert book.sheetnames == ["Özet", "Kategori", "İşletme Projesi"]
    summary = book["Özet"]
    assert summary["A7"].value == "06/2026"
    assert summary["C7"].value == 400
    assert summary["A14"].value == "Toplam"
    assert book["İşletme Projesi"]["B2"].value == 12000


async def test_expense_export_neutralizes_formulas(report: World) -> None:
    await expense(
        report, "123.45", "2026-06-10", description='=HYPERLINK("http://kotu.example","tıkla")'
    )
    book = workbook(await report.get("/expenses/export.xlsx", params={"year": 2026}))
    sheet = book["Giderler"]
    assert [c.value for c in sheet[1]][:6] == [
        "Tarih",
        "Açıklama",
        "Kategori",
        "Tedarikçi",
        "Belge No",
        "Tutar",
    ]
    cell = sheet["B2"]
    assert cell.data_type == "s"  # formül değil, metin
    assert cell.value == '=HYPERLINK("http://kotu.example","tıkla")'
    assert sheet["F2"].value == 123.45 or str(sheet["F2"].value) == "123.45"
    assert sheet["G2"].value == "Hayır"


async def test_statement_export(report: World) -> None:
    await manual(report, "100.00")
    await manual(report, "30.00", "out")
    book = workbook(
        await report.get(f"/cash-accounts/{report.lookups['cash']}/statement/export.xlsx")
    )
    sheet = book["Ekstre"]
    rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
    assert rows[0][1] == "Kasa — devreden"
    assert [r[5] for r in rows[1:3]] == [100, 70]  # eskiden yeniye yürüyen bakiye
    assert (rows[-1][1], rows[-1][5]) == ("Kapanış", 70)


# --- yetki ve izolasyon -----------------------------------------------------------------


async def test_report_permissions_and_isolation(report: World) -> None:
    await expense(report, "400.00", "2026-06-10", paid=True)
    board = await create_user(report.factory, "kurul@test.local")
    await add_site_membership(report.factory, report.site_id, board.id, "Yönetim Kurulu Üyesi")
    board_headers = await login_headers(report.api, "kurul@test.local")
    assert (
        await report.api.get(report.url("/reports/income-expense"), headers=board_headers)
    ).status_code == 200
    tech = await create_user(report.factory, "teknik@test.local")
    await add_site_membership(report.factory, report.site_id, tech.id, "Teknik Personel")
    tech_headers = await login_headers(report.api, "teknik@test.local")
    for path in (
        "/reports/income-expense",
        "/reports/collections",
        "/reports/income-expense/export.xlsx",
    ):
        assert (await report.api.get(report.url(path), headers=tech_headers)).status_code == 403
    other_url = report.url("/reports/income-expense", "yildiz-sitesi")
    response = await report.api.get(other_url, params={"year": 2026}, headers=report.headers)
    other = response.json()
    assert (other["total_expense"], other["total_income"]) == ("0.00", "0.00")
