"""Excel'den bölüm ve kişi aktarımı — dosya, geçici depo, önizleme, yazma (docs/11).

Akış iki adımlıdır: **yükle → önizle** (hiçbir kayıt oluşmaz, dosya geçici depoya yazılır) ve
**onayla → tek transaction ile yaz**. Onaylanmamış dosyalar 6 saat sonra silinir.

Güvenlik (docs/09 §3): uzantı + ZIP imzası, en fazla 5 MB; açılmış boyut ve girdi sayısı
sınırlı (zip bombası); XML defusedxml ile ayrıştırılır (XXE/entity bombası); yalnız hücre
değerleri okunur (formül çalıştırılmaz, makro okunmaz). Geçici dosya adı sunucu üretir;
kullanıcının verdiği ad diske hiç girmez.
"""

import os
import uuid
import warnings
import zipfile
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import and_, select
from sqlalchemy.ext.asyncio import AsyncSession

from site_yonetim.domain.imports.numbers import CellValue
from site_yonetim.domain.imports.unit_validator import (
    Column,
    ImportFileError,
    Issue,
    Severity,
    UnitRow,
    ValidationResult,
    unit_key,
    validate_sheet,
)
from site_yonetim.domain.structure import (
    PartyRole,
    accounts_to_open,
    reference_base,
    reference_code,
)
from site_yonetim.domain.text import tr_lower
from site_yonetim.models import (
    Block,
    LedgerAccount,
    Person,
    Plan,
    Site,
    Unit,
    UnitParty,
    UnitType,
)

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 500
MAX_COLUMNS = 60
PENDING_TTL = timedelta(hours=6)
ZIP_SIGNATURE = b"PK\x03\x04"
SHEET_NAME = "Daireler"
_PENDING_SUFFIX = ".xlsx"
_CLAIM_SUFFIX = ".claim"
# Bozuk arşiv/XML: ParseError SyntaxError'dan, defusedxml'in varlık yasağı ValueError'dan türer.
_PARSE_ERRORS = (zipfile.BadZipFile, KeyError, ValueError, SyntaxError, OSError)


# --- Dosya kontrolü ve okuma ------------------------------------------------------


def check_upload(filename: str | None, data: bytes) -> None:
    """Uzantı, boyut ve içerik imzası. İçerik imzası uzantıdan bağımsız doğrulanır."""
    if not (filename or "").lower().endswith(".xlsx"):
        raise ImportFileError(
            "invalid_file_type", "Yalnız .xlsx uzantılı Excel dosyası yüklenebilir."
        )
    if len(data) > MAX_UPLOAD_BYTES:
        raise ImportFileError("file_too_large", "Dosya en fazla 5 MB olabilir.")
    if not data.startswith(ZIP_SIGNATURE):
        raise ImportFileError(
            "invalid_file_content",
            "Dosya geçerli bir Excel (.xlsx) dosyası değil. Excel'de 'Farklı Kaydet → "
            "Excel Çalışma Kitabı (.xlsx)' ile kaydedip yeniden deneyin.",
        )


def _check_archive(data: bytes) -> None:
    """Zip bombasına karşı: girdi sayısı ve **beyan edilen** açılmış boyut sınırı.

    `zipfile` okurken beyan edilen boyutun ötesine geçmez; beyanı büyük olan reddedilir.
    """
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise _unreadable() from exc
    if len(entries) > MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > (
        MAX_UNCOMPRESSED_BYTES
    ):
        raise ImportFileError(
            "file_too_complex",
            "Dosya açıldığında çok büyük; lütfen yalnız daire listesini aktarın.",
        )
    if not any(e.filename == "xl/workbook.xml" for e in entries):
        raise _unreadable()


def _unreadable() -> ImportFileError:
    return ImportFileError(
        "invalid_file_content",
        "Dosya okunamadı. Şablonu indirip verileri ona kopyalayarak yeniden deneyin.",
    )


@contextmanager
def open_sheet(data: bytes) -> Iterator[Iterator[tuple[CellValue, ...]]]:
    """Aktarım sayfasının satırları (`Daireler`, yoksa ilk sayfa), yalnız değerler."""
    _check_archive(data)
    with warnings.catch_warnings():
        # openpyxl desteklemediği uzantılar için uyarı basar (veri doğrulama, koşullu biçim).
        warnings.simplefilter("ignore")
        try:
            workbook = openpyxl.load_workbook(
                BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
        except (InvalidFileException, *_PARSE_ERRORS) as exc:
            raise _unreadable() from exc
    try:
        names = workbook.sheetnames
        sheet = workbook[SHEET_NAME] if SHEET_NAME in names else workbook.worksheets[0]
        yield sheet.iter_rows(max_col=MAX_COLUMNS, values_only=True)  # type: ignore[misc]
    finally:
        workbook.close()


def read_and_validate(data: bytes) -> ValidationResult:
    """CPU'ya bağlı ve senkron — uç nokta bunu iş parçacığında çağırır."""
    with open_sheet(data) as rows:
        try:
            return validate_sheet(rows)
        except _PARSE_ERRORS as exc:
            # Satırlar okundukça ayrıştırılır; bozuk sayfa XML'i burada ortaya çıkar.
            raise _unreadable() from exc


# --- Şablon (docs/11 §2) ----------------------------------------------------------

_EXAMPLE_ROWS: tuple[tuple[object, ...], ...] = (
    ("A", "1", 1, "2+1", 104.5, 86, 40, 10000, "Konut", "Ayşe", "Yılmaz", "5321234567",
     "ayse@ornek.com", None, None, None, None),
    ("A", "2", 1, "1+1", 68, 56.4, 26, 10000, "Konut", "Mehmet", "Kaya", "05339876543",
     None, "Elif", "Demir", "5445556677", None),
    ("B", "Z1", 0, "Dükkan", 120, 110, 54, 10000, "Ticari", "Ali", "Çelik", "+905367778899",
     None, None, None, None, None),
)  # fmt: skip

_HELP_LINES = (
    ("Excel'den daire ve sakin aktarımı", True),
    ("", False),
    ("Verileri 'Daireler' sayfasına yazın; örnek satırları silin.", False),
    ("Sütun sırası önemli değil, başlık adına bakılır.", False),
    ("Zorunlu: Daire No, Malik Ad, Malik Soyad. Diğerleri boş bırakılabilir.", False),
    ("Arsa payı: pay ve payda birlikte girilir (ör. 40 / 10000).", False),
    ("Kullanım: Konut, Ticari (dükkan, ofis), Depo, Otopark. Boşsa konut.", False),
    ("Telefon: 5 ile başlayan 10 haneli cep numarası (0532…, +90… de olur).", False),
    ("Metrekare: 104,5 ya da 104.5 yazılabilir.", False),
    ("Kiracı yoksa kiracı sütunlarını boş bırakın.", False),
    ("", False),
    ("Yükledikten sonra önizleme gösterilir; onaylamadan hiçbir kayıt oluşmaz.", False),
    ("Sistemde zaten kayıtlı bölümler (aynı blok + numara) atlanır, üzerine yazılmaz.", False),
    ("En fazla 5 MB ve 5.000 satır.", False),
)


def build_template() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.worksheets[0]  # yeni çalışma kitabında her zaman bir sayfa var
    sheet.title = SHEET_NAME
    headers = [c.value for c in Column]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for example in _EXAMPLE_ROWS:
        sheet.append(list(example))
    sheet.freeze_panes = "A2"
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(12, len(header) + 4)
    # Telefon ve daire numarası metin olarak kalsın (Excel baştaki 0'ı silmesin).
    for column in (Column.NUMBER, Column.OWNER_PHONE, Column.TENANT_PHONE):
        letter = get_column_letter(headers.index(column.value) + 1)
        sheet.column_dimensions[letter].number_format = "@"

    help_sheet = workbook.create_sheet("Açıklama")
    for text, bold in _HELP_LINES:
        help_sheet.append([text])
        if bold:
            help_sheet.cell(row=help_sheet.max_row, column=1).font = Font(bold=True, size=13)
    help_sheet.column_dimensions["A"].width = 90

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- Geçici depo --------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class PendingImport:
    import_id: uuid.UUID
    expires_at: datetime


class ImportStore:
    """Onay bekleyen dosyalar: `{kök}/{site_id}/{user_id}/{import_id}.xlsx`.

    Yol bileşenlerinin hepsi sunucunun ürettiği UUID'dir. Bir aktarımı yalnız onu yükleyen
    kullanıcı, aynı sitede onaylayabilir. Birden çok API kopyası çalışıyorsa kök ortak bir
    birim olmalı (docs/11 §1).
    """

    def __init__(self, root: Path) -> None:
        self.root = root

    def _dir(self, site_id: uuid.UUID, user_id: uuid.UUID) -> Path:
        return self.root / site_id.hex / user_id.hex

    def save(
        self, site_id: uuid.UUID, user_id: uuid.UUID, data: bytes, now: datetime
    ) -> PendingImport:
        self.purge_expired(now)
        folder = self._dir(site_id, user_id)
        folder.mkdir(mode=0o700, parents=True, exist_ok=True)
        import_id = uuid.uuid4()  # tahmin edilemez
        path = folder / f"{import_id.hex}{_PENDING_SUFFIX}"
        fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, "O_BINARY", 0), 0o600)
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
        stamp = now.timestamp()
        os.utime(path, (stamp, stamp))
        return PendingImport(import_id, now + PENDING_TTL)

    @contextmanager
    def claim(
        self, site_id: uuid.UUID, user_id: uuid.UUID, import_id: uuid.UUID, now: datetime
    ) -> Iterator[bytes | None]:
        """Dosyayı atomik olarak sahiplenir: aynı aktarım iki kez onaylanamaz.

        Blok hatasız biterse dosya silinir; hata olursa geri konur (kullanıcı yeniden dener).
        Dosya yoksa, süresi dolmuşsa ya da başkasınınsa `None` verir.
        """
        folder = self._dir(site_id, user_id)
        pending = folder / f"{import_id.hex}{_PENDING_SUFFIX}"
        claimed = folder / f"{import_id.hex}.{uuid.uuid4().hex}{_CLAIM_SUFFIX}"
        try:
            pending.rename(claimed)
        except FileNotFoundError:
            yield None
            return
        if self._expired(claimed, now):
            claimed.unlink(missing_ok=True)
            yield None
            return
        try:
            yield claimed.read_bytes()
        except BaseException:
            claimed.rename(pending)
            raise
        claimed.unlink(missing_ok=True)

    def purge_expired(self, now: datetime) -> int:
        """Süresi dolmuş (onaylanmamış ya da yarım kalmış) dosyaları siler."""
        if not self.root.is_dir():
            return 0
        removed = 0
        for path in self.root.glob("*/*/*"):
            if path.is_file() and self._expired(path, now):
                path.unlink(missing_ok=True)
                removed += 1
        return removed

    @staticmethod
    def _expired(path: Path, now: datetime) -> bool:
        try:
            modified = datetime.fromtimestamp(path.stat().st_mtime, UTC)
        except FileNotFoundError:
            return True
        return now - modified >= PENDING_TTL


# --- Önizleme -----------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class Preview:
    result: ValidationResult
    existing: frozenset[tuple[str, str]]  # sitede zaten kayıtlı bölümler — atlanacak
    notices: list[Issue] = field(default_factory=list)  # dosya geneli uyarılar

    @property
    def new_rows(self) -> list[UnitRow]:
        return [r for r in self.result.rows if r.key not in self.existing]


async def _existing_unit_keys(session: AsyncSession) -> frozenset[tuple[str, str]]:
    rows = await session.execute(
        select(Block.name, Unit.number).join(
            Block, and_(Block.id == Unit.block_id, Block.site_id == Unit.site_id)
        )
    )
    return frozenset(unit_key(block, number) for block, number in rows)


async def build_preview(session: AsyncSession, site: Site, result: ValidationResult) -> Preview:
    existing = await _existing_unit_keys(session)
    preview = Preview(result, existing)
    # Öneri (docs/11 §4): plan tavanı aşılacaksa uyar, engelleme.
    max_units = await session.scalar(select(Plan.max_units).where(Plan.id == site.plan_id))
    if max_units is not None:
        after = len(existing) + len(preview.new_rows)
        if after > max_units:
            preview.notices.append(
                Issue(
                    None,
                    None,
                    f"Aktarımdan sonra bölüm sayısı {after} olacak; planınızın sınırı "
                    f"{max_units}. Aktarım yapılır, ancak plan yükseltmesi için platform "
                    "yöneticinizle görüşün.",
                    Severity.WARNING,
                )
            )
    return preview


# --- Yazma (docs/11 §4) -------------------------------------------------------------


@dataclass(slots=True)
class ImportOutcome:
    created_units: int = 0
    created_people: int = 0
    created_accounts: int = 0
    created_blocks: int = 0
    created_unit_types: int = 0
    skipped: list[str] = field(default_factory=list)

    @property
    def message(self) -> str:
        if self.created_units:
            text = f"{self.created_units} bölüm ve {self.created_people} kişi aktarıldı."
        else:
            text = "Yeni bölüm aktarılmadı."
        if self.skipped:
            text += f" {len(self.skipped)} bölüm zaten kayıtlı olduğu için atlandı."
        return text


def _name_key(name: str) -> str:
    return tr_lower(" ".join(name.split()))


def import_start_date(today: date) -> date:
    """Taraf başlangıcı: içinde bulunulan yılın 1 Ocak'ı (docs/11 §4)."""
    return date(today.year, 1, 1)


async def apply_import(
    session: AsyncSession, rows: list[UnitRow], *, start_date: date
) -> ImportOutcome:
    """Önizlenen satırları yazar. Transaction'ı çağıran yönetir (hepsi ya da hiçbiri).

    Sorgu sayısı satır sayısından bağımsızdır: mevcut bloklar, tipler, bölümler ve hesap
    kodları bir kez okunur, yeni kayıtlar tablo sırasıyla toplu yazılır.
    """
    outcome = ImportOutcome()
    blocks = {_name_key(b.name): b for b in await session.scalars(select(Block))}
    unit_types = {_name_key(t.name): t for t in await session.scalars(select(UnitType))}
    block_order = max((b.sort_order for b in blocks.values()), default=-1)
    type_order = max((t.sort_order for t in unit_types.values()), default=-1)
    existing = await _existing_unit_keys(session)
    taken_codes = set(await session.scalars(select(LedgerAccount.reference_code)))

    new_units: list[tuple[UnitRow, Unit, Block]] = []
    for row in rows:
        if row.key in existing:
            outcome.skipped.append(row.display_name)
            continue
        block = blocks.get(_name_key(row.block))
        if block is None:
            block_order += 1
            block = Block(id=uuid.uuid7(), name=row.block, sort_order=block_order)
            blocks[_name_key(row.block)] = block
            session.add(block)
            outcome.created_blocks += 1
        unit_type = None
        if row.unit_type:
            unit_type = unit_types.get(_name_key(row.unit_type))
            if unit_type is None:
                type_order += 1
                # Ağırlık 1: yönetici sonra düzeltir (docs/11 §4).
                unit_type = UnitType(
                    id=uuid.uuid7(), name=row.unit_type, weight=Decimal(1), sort_order=type_order
                )
                unit_types[_name_key(row.unit_type)] = unit_type
                session.add(unit_type)
                outcome.created_unit_types += 1
        unit = Unit(
            id=uuid.uuid7(),
            block_id=block.id,
            number=row.number,
            floor=row.floor,
            unit_type_id=unit_type.id if unit_type else None,
            gross_area=row.gross_area,
            net_area=row.net_area,
            land_share_numerator=row.land_share_numerator,
            land_share_denominator=row.land_share_denominator,
            usage=row.usage.value,
        )
        new_units.append((row, unit, block))
    if not new_units:
        return outcome
    await session.flush()  # bloklar, tipler
    session.add_all(unit for _, unit, _ in new_units)
    await session.flush()

    parties: list[tuple[Unit, Block, Person, PartyRole, bool]] = []
    for row, unit, block in new_units:
        has_tenant = row.tenant is not None
        people = [(row.owner, PartyRole.OWNER)]
        if row.tenant is not None:
            people.append((row.tenant, PartyRole.TENANT))
        for data, role in people:
            person = Person(
                id=uuid.uuid7(),
                first_name=data.first_name,
                last_name=data.last_name,
                phone=data.phone,
                email=data.email,
            )
            session.add(person)
            parties.append((unit, block, person, role, has_tenant))
    await session.flush()

    for unit, block, person, role, has_tenant in parties:
        session.add(
            UnitParty(unit_id=unit.id, person_id=person.id, role=role.value, start_date=start_date)
        )
        base = reference_base(block.name, unit.number)
        for kind, letter in accounts_to_open(role, unit_has_active_tenant=has_tenant):
            code = reference_code(base, letter, taken_codes)
            taken_codes.add(code)
            session.add(
                LedgerAccount(
                    unit_id=unit.id, person_id=person.id, kind=kind.value, reference_code=code
                )
            )
            outcome.created_accounts += 1
    await session.flush()

    outcome.created_units = len(new_units)
    outcome.created_people = len(parties)
    return outcome
