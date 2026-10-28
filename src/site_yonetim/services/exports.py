"""Excel çıktıları — docs/06 §1.6, §2.7–§2.9.

Güvenlik: kullanıcıdan gelen metin hücreye **her zaman metin** olarak yazılır — `=`, `+`, `-`,
`@` ile başlayan bir açıklama formül olarak çalışmaz (formül enjeksiyonu). Para `Decimal` olarak
yazılır, `#,##0.00` biçimiyle gösterilir (float yok).
"""

import datetime as dt
from collections.abc import Sequence
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from site_yonetim.domain.text import format_period_tr

MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "dd.mm.yyyy"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

type CellValue = str | int | Decimal | dt.date | bool | None


def _write_row(sheet: Worksheet, values: Sequence[CellValue], *, bold: bool = False) -> None:
    row = sheet.max_row + 1 if sheet.max_row > 1 or sheet["A1"].value is not None else 1
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column)
        if isinstance(value, str):
            cell.value = value
            cell.data_type = "s"  # formül değil: metin
        elif isinstance(value, bool):
            cell.value = "Evet" if value else "Hayır"
            cell.data_type = "s"
        else:
            cell.value = value
            if isinstance(value, Decimal):
                cell.number_format = MONEY_FORMAT
            elif isinstance(value, dt.date):
                cell.number_format = DATE_FORMAT
        if bold:
            cell.font = Font(bold=True)


def _sheet(
    workbook: Workbook, title: str, header: Sequence[str], *, first: bool = False
) -> Worksheet:
    sheet = workbook.worksheets[0] if first else workbook.create_sheet()
    sheet.title = title
    _write_row(sheet, header, bold=True)
    sheet.freeze_panes = "A2"
    for index, name in enumerate(header, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(12, len(name) + 4)
    return sheet


def _bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- Giderler --------------------------------------------------------------------


EXPENSE_HEADER = (
    "Tarih", "Açıklama", "Kategori", "Tedarikçi", "Belge No", "Tutar", "Ödendi",
    "Ödeme Tarihi", "Hesap", "Durum",
)  # fmt: skip


def expense_rows_xlsx(rows: Sequence[Sequence[CellValue]]) -> bytes:
    workbook = Workbook()
    sheet = _sheet(workbook, "Giderler", EXPENSE_HEADER, first=True)
    sheet.column_dimensions["B"].width = 40
    for row in rows:
        _write_row(sheet, row)
    return _bytes(workbook)


# --- Kasa ekstresi -------------------------------------------------------------------


STATEMENT_HEADER = ("Tarih", "Açıklama", "Referans", "Giren", "Çıkan", "Bakiye", "Kaynak", "Not")


def statement_xlsx(
    account_name: str,
    opening: Decimal,
    closing: Decimal,
    rows: Sequence[Sequence[CellValue]],
) -> bytes:
    """Sayfalama yok, tüm aralık; eskiden yeniye (yürüyen bakiye okunur kalsın)."""
    workbook = Workbook()
    sheet = _sheet(workbook, "Ekstre", STATEMENT_HEADER, first=True)
    sheet.column_dimensions["B"].width = 45
    _write_row(sheet, (None, f"{account_name} — devreden", None, None, None, opening))
    for row in rows:
        _write_row(sheet, row)
    _write_row(sheet, (None, "Kapanış", None, None, None, closing), bold=True)
    return _bytes(workbook)


# --- Gelir–gider raporu ----------------------------------------------------------------


def report_xlsx(
    year: int,
    months: Sequence[tuple[int, Decimal, Decimal, Decimal]],
    categories: Sequence[tuple[str, Decimal, int, Decimal]],
    budget: Sequence[tuple[str, Decimal, Decimal, Decimal, Decimal, bool]],
) -> bytes:
    """3 sayfa: Özet (ay ay), Kategori, İşletme Projesi (docs/06 §2.9)."""
    workbook = Workbook()
    summary = _sheet(workbook, "Özet", ("Dönem", "Gelir", "Gider", "Fark"), first=True)
    for month, income, expense, difference in months:
        _write_row(summary, (format_period_tr(year, month), income, expense, difference))
    _write_row(
        summary,
        (
            "Toplam",
            sum((m[1] for m in months), Decimal(0)),
            sum((m[2] for m in months), Decimal(0)),
            sum((m[3] for m in months), Decimal(0)),
        ),
        bold=True,
    )
    category = _sheet(workbook, "Kategori", ("Kategori", "Tutar", "Kayıt Sayısı", "Pay (%)"))
    for name, amount, count, share in categories:
        _write_row(category, (name, amount, count, share))
    plan = _sheet(
        workbook,
        "İşletme Projesi",
        ("Kategori", "Bütçelenen", "Gerçekleşen", "Fark", "Kullanım (%)", "Aşıldı"),
    )
    for name, budgeted, actual, difference, usage, is_over in budget:
        _write_row(plan, (name, budgeted, actual, difference, usage, is_over))
    return _bytes(workbook)


# --- Katalog raporu ve envanter (RAP-01, RAP-04) ----------------------------------------------

_SHEET_FORBIDDEN = str.maketrans(dict.fromkeys("[]:*?/\\", " "))


def _sheet_title(title: str, used: set[str]) -> str:
    base = title.translate(_SHEET_FORBIDDEN).strip()[:28] or "Sayfa"
    name, n = base, 2
    while name in used:
        name, n = f"{base} {n}", n + 1
    used.add(name)
    return name


def content_xlsx(
    title: str,
    period: str,
    summary: Sequence[tuple[str, CellValue]],
    tables: Sequence[tuple[str, Sequence[str], Sequence[Sequence[CellValue]]]],
    metrics: Sequence[tuple[str, str]],
) -> bytes:
    """Katalog raporu: Özet (başlık, dönem, göstergeler, metrik tanımları) + her tablo bir sayfa.
    Planlı rapor gönderimi de bunu kullanır."""
    workbook = Workbook()
    used = {"Özet"}
    sheet = _sheet(workbook, "Özet", (title, period), first=True)
    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["B"].width = 24
    for label, value in summary:
        _write_row(sheet, (label, value))
    if metrics:
        _write_row(sheet, ())
        _write_row(sheet, ("Metrik", "Nasıl hesaplanır"), bold=True)
        for name, definition in metrics:
            _write_row(sheet, (name, definition))
    for table_title, header, rows in tables:
        page = _sheet(workbook, _sheet_title(table_title, used), header)
        for row in rows:
            _write_row(page, row)
    return _bytes(workbook)


_CSV_RISKY = ("=", "+", "-", "@", "\t", "\r")


def _csv_cell(value: CellValue) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Evet" if value else "Hayır"
    if isinstance(value, dt.date):
        return value.strftime("%d.%m.%Y")
    text = str(value)
    if isinstance(value, str) and text.startswith(_CSV_RISKY):
        text = "'" + text  # formül enjeksiyonu
    if any(c in text for c in ';"\n'):
        text = '"' + text.replace('"', '""') + '"'
    return text


BOM = chr(0xFEFF)
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"


def rows_csv(header: Sequence[str], rows: Sequence[Sequence[CellValue]]) -> bytes:
    """Excel'in Türkçe ayarıyla açılan CSV: `;` ayraç, UTF-8 BOM, formül enjeksiyonuna karşı
    `=`/`+`/`-`/`@` ile başlayan metin `'` ile başlar."""
    lines = [";".join(_csv_cell(h) for h in header)]
    lines += [";".join(_csv_cell(v) for v in row) for row in rows]
    return (BOM + "\r\n".join(lines) + "\r\n").encode("utf-8")


def rows_xlsx(
    sheet_title: str, header: Sequence[str], rows: Sequence[Sequence[CellValue]]
) -> bytes:
    workbook = Workbook()
    sheet = _sheet(workbook, sheet_title, header, first=True)
    for row in rows:
        _write_row(sheet, row)
    return _bytes(workbook)
