"""Muhasebe aktarımı (Mod P) — saf kurallar (Yazılım Dokümanı FIN-14, §8.4; servis isteği 32).

Platform genel bir muhasebe fişi dosyası üretir; muhasebe sorumlusu dosyayı muhasebe programına
yükler. Kivi'ye özel biçim **K7 kararına (Kivi görüşmesi) bağlı** — dosya
`AccountingExportAdapter` arayüzüyle yazılır, Kivi/Logo/Mikro biçimi aynı arayüze eklenir.

- Her fiş tek borç + tek alacak satırıdır; dosyada borç toplamı = alacak toplamı.
- Her fiş bir kez aktarılır ve kilitlenir; düzeltme ters fişle (sonraki dosyada).
- Ekrandaki varsayılan hesap kodları yalnız öneridir — mali müşavir teyidi gerekir.
"""

import csv
import io
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import Protocol

from openpyxl import Workbook

from site_yonetim.domain.text import format_date_tr


@dataclass(frozen=True, slots=True)
class MappingDef:
    key: str
    label: str
    debit_account: str
    credit_account: str


# Önerilen tek düzen hesap planı kodları (mali müşavir teyidi gerekir).
DEFAULT_MAPPINGS: tuple[MappingDef, ...] = (
    MappingDef("charge", "Aidat tahakkuku", "120", "600"),
    MappingDef("ad_hoc_charge", "Dönem dışı ücret", "120", "602"),
    MappingDef("payment_cash", "Tahsilat — nakit", "100", "120"),
    MappingDef("payment_bank_transfer", "Tahsilat — havale/EFT", "102", "120"),
    MappingDef("payment_credit_card", "Tahsilat — kredi kartı", "108", "120"),
    MappingDef("payment_other", "Tahsilat — diğer", "", "120"),
    MappingDef("expense", "Gider ödemesi", "770", "102"),
    MappingDef("refund", "İade", "120", "102"),
)
MAPPING_KEYS = tuple(m.key for m in DEFAULT_MAPPINGS)
LABELS = {m.key: m.label for m in DEFAULT_MAPPINGS}
_CODE = re.compile(r"^[0-9][0-9.]{0,19}$")


def check_code(value: str) -> bool:
    """Hesap kodu: rakamla başlar, rakam ve nokta (ör. 120.01); boş = eşlenmemiş."""
    return value == "" or bool(_CODE.match(value))


@dataclass(frozen=True, slots=True)
class Voucher:
    key: str  # "kaynak:kimlik" — fişin tekil anahtarı (bir kez aktarılır)
    source: str  # eşleme anahtarı (charge, payment_cash…)
    source_id: str
    date: date
    description: str
    debit_account: str
    credit_account: str
    amount: Decimal


class AccountingExportAdapter(Protocol):
    """Fiş dosyası biçimi. Kivi/Logo/Mikro biçimleri K7 kararından sonra eklenir."""

    extension: str
    media_type: str

    def render(self, vouchers: Sequence[Voucher]) -> bytes: ...


HEADERS = ("Fiş tarihi", "Kaynak", "Kaynak no", "Açıklama", "Borç hesabı", "Alacak hesabı", "Tutar")


def _amount_tr(value: Decimal) -> str:
    return f"{value:.2f}".replace(".", ",")


def _safe(text: str) -> str:
    """Hesap tablosu formül enjeksiyonu: =,+,-,@ ile başlayan metin tırnaklanır."""
    return f"'{text}" if text[:1] in ("=", "+", "-", "@") else text


class GenericCsvAdapter:
    """Genel CSV: `;` ayraçlı, UTF-8 BOM'lu, tarih `15.09.2026`, tutar `1234,56`."""

    extension = "csv"
    media_type = "text/csv; charset=utf-8"

    def render(self, vouchers: Sequence[Voucher]) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
        writer.writerow(HEADERS)
        for v in vouchers:
            writer.writerow(
                [
                    format_date_tr(v.date),
                    LABELS.get(v.source, v.source),
                    v.source_id,
                    _safe(v.description),
                    v.debit_account,
                    v.credit_account,
                    _amount_tr(v.amount),
                ]
            )
        return ("﻿" + buffer.getvalue()).encode("utf-8")


class GenericXlsxAdapter:
    """Aynı sütunlarla Excel; tutar sayı hücresi (Decimal), tarih tarih hücresi."""

    extension = "xlsx"
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def render(self, vouchers: Sequence[Voucher]) -> bytes:
        book = Workbook()
        sheet = book.create_sheet("Fişler", 0)
        for extra in book.sheetnames[1:]:
            del book[extra]
        sheet.append(list(HEADERS))
        for v in vouchers:
            sheet.append(
                [
                    v.date,
                    LABELS.get(v.source, v.source),
                    v.source_id,
                    _safe(v.description),
                    v.debit_account,
                    v.credit_account,
                    v.amount,
                ]
            )
        out = io.BytesIO()
        book.save(out)
        return out.getvalue()


ADAPTERS: dict[str, AccountingExportAdapter] = {
    "csv": GenericCsvAdapter(),
    "xlsx": GenericXlsxAdapter(),
}
